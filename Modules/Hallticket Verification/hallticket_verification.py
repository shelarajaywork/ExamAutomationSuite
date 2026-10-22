"""
Hall Ticket PDF -> Excel extractor.

Run with:  streamlit run app.py
"""

import io
import re
from datetime import datetime

import pandas as pd
import pdfplumber
import streamlit as st

# --------------------------------------------------------------------------
# Parsing helpers
# --------------------------------------------------------------------------

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"\b[6-9]\d{9}\b")
CITY_PIN_RE = re.compile(r"([A-Za-z .()'-]+?)\s*-\s*(\d{6})")

HEADER_LABELS = [
    "Student Number",
    "Roll No.",
    "Gender",
    "Student Name",
    "Program Name",
    "Month & Year of\nExamination",
    "Academic Year",
    "Academic Session",
    "Address",
]

# Columns from Students sheet to exclude in Exam Schedule sheet
EXCLUDE_STUDENT_COLS_IN_SCHEDULE = {
    "Address",
    "City",
    "Pincode",
    "Mobile 1",
    "Mobile 2",
    "Email",
    "Subjects",
    "Page",
    "Source File",
}


def _clean(value):
    """Collapse the newlines pdfplumber leaves inside wrapped cells."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _value_after(rows, label):
    """Return the first non-empty cell to the right of `label` in any row."""
    for row in rows:
        for idx, cell in enumerate(row):
            if _clean(cell).lower() == label.replace("\n", " ").lower():
                for nxt in row[idx + 1:]:
                    val = _clean(nxt)
                    if val and val.lower() not in ("photo", "signature"):
                        return val
    return ""


def _flags(rows):
    """Pick up standalone markers such as PWD that sit in the header row."""
    found = []
    for row in rows:
        for cell in row:
            val = _clean(cell)
            if val.upper() in ("PWD", "PH", "DIVYANG") and val.upper() not in found:
                found.append(val.upper())
    return ", ".join(found)


def split_address(raw):
    """Break the single address blob into address / city / pin / phones / email."""
    raw = _clean(raw)
    email = ""
    match = EMAIL_RE.search(raw)
    if match:
        email = match.group(0)
        raw = raw.replace(email, " ")

    phones = []
    for phone in PHONE_RE.findall(raw):
        if phone not in phones:
            phones.append(phone)
    for phone in phones:
        raw = raw.replace(phone, " ")

    city, pin = "", ""
    city_match = CITY_PIN_RE.search(raw)
    if city_match:
        city = city_match.group(1).strip(" ,-")
        pin = city_match.group(2)
        raw = raw[: city_match.start()] + " " + raw[city_match.end():]

    parts = [p.strip(" ,") for p in raw.split(",")]
    address = ", ".join(p for p in parts if p)
    address = re.sub(r"\s+", " ", address).strip(" ,")

    return {
        "Address": address,
        "City": city,
        "Pincode": pin,
        "Mobile 1": phones[0] if phones else "",
        "Mobile 2": phones[1] if len(phones) > 1 else "",
        "Email": email,
    }


def parse_date(text):
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_page(page):
    """Return (student_dict, [subject_dict, ...]) for one hall ticket page."""
    tables = page.extract_tables()
    if not tables:
        return None, []

    header_rows, subject_rows = [], []
    for table in tables:
        flat = " ".join(_clean(c) for row in table for c in row)
        if "Student Number" in flat or "Student Name" in flat:
            header_rows.extend(table)
        elif "Module Code" in flat or re.search(r"\b[A-Z]{4,6}\d{3}\b", flat):
            subject_rows.extend(table)

    student = {
        "Student Number": _value_after(header_rows, "Student Number"),
        "Roll No.": _value_after(header_rows, "Roll No."),
        "Student Name": _value_after(header_rows, "Student Name"),
        "Gender": _value_after(header_rows, "Gender"),
        "Program Name": _value_after(header_rows, "Program Name"),
        "Exam Month & Year": _value_after(header_rows, "Month & Year of Examination"),
        "Academic Year": _value_after(header_rows, "Academic Year"),
        "Academic Session": _value_after(header_rows, "Academic Session"),
        "Category": _flags(header_rows),
    }
    student.update(split_address(_value_after(header_rows, "Address")))

    subjects = []
    for row in subject_rows:
        cells = [_clean(c) for c in row]
        if len(cells) < 6 or not re.fullmatch(r"\d{1,2}", cells[0]):
            continue  # header row or stray line
        date_text = cells[3]
        time_text = cells[4]
        start, _, end = time_text.partition("-")
        subjects.append(
            {
                "Sr. No.": int(cells[0]),
                "Module Code": cells[1],
                "Module Description": cells[2],
                "Exam Date": parse_date(date_text) or date_text,
                "Start Time": start.strip(),
                "End Time": end.strip(),
                "Room": cells[5],
            }
        )

    subjects.sort(key=lambda s: s["Sr. No."])
    return student, subjects


def parse_pdf(file_obj, source_name=""):
    students, schedule = [], []
    with pdfplumber.open(file_obj) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            student, subjects = parse_page(page)
            if not student or not student.get("Student Number"):
                continue
            student["Subjects"] = len(subjects)
            student["Page"] = page_no
            if source_name:
                student["Source File"] = source_name
            students.append(student)

            # Extract student columns required in Exam Schedule (excluding address/meta fields)
            student_schedule_details = {
                k: v
                for k, v in student.items()
                if k not in EXCLUDE_STUDENT_COLS_IN_SCHEDULE
            }

            for subject in subjects:
                schedule.append(
                    {
                        **student_schedule_details,
                        **subject,
                    }
                )
    return pd.DataFrame(students), pd.DataFrame(schedule)


def build_timetable(schedule_df):
    """Distinct exam slots across all students, i.e. the datesheet itself."""
    if schedule_df.empty:
        return pd.DataFrame()
    cols = ["Module Code", "Module Description", "Exam Date", "Start Time", "End Time"]
    table = (
        schedule_df.groupby(cols, dropna=False)
        .agg(Students=("Student Number", "nunique"),
             Rooms=("Room", lambda s: ", ".join(sorted(set(s)))))
        .reset_index()
        .sort_values(["Exam Date", "Start Time"])
    )
    return table


CONTEXT_COLS = [
    "Program Name",
    "Exam Month & Year",
    "Academic Year",
    "Academic Session",
]
SLOT_COLS = ["Exam Date", "Start Time", "End Time", "Module Code", "Module Description"]

# Desired leading columns in the pivot (everything else = room columns, appended sorted)
PIVOT_LEAD_COLS = [
    "Exam Date", "Start Time", "End Time",
    "Program Name", "Academic Session",
    "Module Code", "Module Description",
    "Grand Total", "PwD Students",
]


def _room_sort_key(room_name):
    """Natural sort key: numeric suffixes first (as int), then alpha."""
    import re as _re
    m = _re.search(r'[\s.]+([A-Z0-9]+)\s*$', str(room_name))
    suffix = m.group(1) if m else room_name
    try:
        return (0, int(suffix), "")      # numeric rooms first, in numeric order
    except ValueError:
        return (1, 0, suffix)            # alpha rooms after, alphabetically


def _room_short_label(room_name):
    """Return just the trailing room number / letter from the full room string."""
    import re as _re
    m = _re.search(r'[\s.]+([A-Z0-9]+)\s*$', str(room_name))
    return m.group(1) if m else room_name


def build_pivot(students_df, schedule_df):
    """Students per room for every exam slot, split by program / session."""
    if schedule_df.empty or students_df.empty:
        return pd.DataFrame(), [], {}

    keep = ["Student Number", "Category"] + CONTEXT_COLS

    # Avoid duplicate column name collision (_x, _y) when merging
    cols_to_merge = [
        c
        for c in keep
        if c in students_df.columns
        and (c == "Student Number" or c not in schedule_df.columns)
    ]
    if len(cols_to_merge) > 1:
        merged = schedule_df.merge(
            students_df[cols_to_merge].drop_duplicates("Student Number"),
            on="Student Number",
            how="left",
        )
    else:
        merged = schedule_df.copy()

    for col in CONTEXT_COLS:
        if col not in merged.columns:
            merged[col] = ""
        merged[col] = merged[col].fillna("")
    merged["Category"] = merged.get("Category", "").fillna("")

    group = CONTEXT_COLS + SLOT_COLS

    counts = (
        merged.pivot_table(
            index=group,
            columns="Room",
            values="Student Number",
            aggfunc="nunique",
            fill_value=0,
        )
        .reset_index()
    )
    raw_room_cols = [c for c in counts.columns if c not in group]

    # Sort room columns with natural / numeric sort
    raw_room_cols_sorted = sorted(raw_room_cols, key=_room_sort_key)

    # Build short-label rename map  {full_name: short_label}
    # Guarantee uniqueness: if two full names collapse to the same short label, keep full
    short_labels = {}
    seen = {}
    for full in raw_room_cols_sorted:
        short = _room_short_label(full)
        if short in seen:
            # collision: revert both to full name
            short_labels[seen[short]] = seen[short]
            short_labels[full] = full
        else:
            seen[short] = full
            short_labels[full] = short

    # Rename room columns in counts df
    counts = counts.rename(columns=short_labels)
    short_room_cols = [short_labels[r] for r in raw_room_cols_sorted]

    pwd = (
        merged[merged["Category"].str.upper().str.contains("PWD", na=False)]
        .groupby(group, dropna=False)["Student Number"]
        .nunique()
        .rename("PwD Students")
        .reset_index()
    )
    pivot = counts.merge(pwd, on=group, how="left")
    pivot["PwD Students"] = pivot["PwD Students"].fillna(0).astype(int)
    pivot["Grand Total"] = pivot[[short_labels[r] for r in raw_room_cols_sorted]].sum(axis=1).astype(int)

    # Reorder columns: lead cols then room cols
    all_cols = PIVOT_LEAD_COLS + short_room_cols
    # keep only cols that actually exist
    all_cols = [c for c in all_cols if c in pivot.columns]
    pivot = pivot[all_cols].sort_values(
        ["Exam Date", "Start Time", "Program Name", "Academic Session"]
    ).reset_index(drop=True)

    return pivot, short_room_cols, short_labels


def write_pivot_sheet(worksheet, pivot, room_cols):
    """Lay the pivot out with the requested column order and formatting."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    # The pivot df already carries columns in the correct order (set by build_pivot).
    headers = list(pivot.columns)
    n_cols = len(headers)

    # Identify semantic groups by name
    centre_always = {"Exam Date", "Start Time", "End Time",
                     "Grand Total", "PwD Students"} | set(room_cols)
    num_cols = {"Grand Total", "PwD Students"} | set(room_cols)

    pwd_idx   = headers.index("PwD Students") + 1   # 1-based
    total_idx = headers.index("Grand Total")  + 1

    # Index of first room column (1-based)
    first_room_idx = headers.index(room_cols[0]) + 1 if room_cols else n_cols + 1

    thin   = Side(style="thin",   color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_fill  = PatternFill("solid", fgColor="D9E1F2")   # blue-grey
    pwd_fill     = PatternFill("solid", fgColor="DDEBF7")   # soft blue
    total_fill   = PatternFill("solid", fgColor="D6E4BC")   # soft green for Grand Total
    room_fill    = PatternFill("solid", fgColor="FFF2CC")   # soft amber for rooms
    footer_fill  = PatternFill("solid", fgColor="F2F2F2")   # light grey for total row

    centre = Alignment(horizontal="center", vertical="center")

    # ── Row 1: title ─────────────────────────────────────────────────────────
    cell = worksheet.cell(row=1, column=1,
                          value="Count of Students by Exam Slot and Room")
    cell.font = Font(bold=True, size=13)
    worksheet.merge_cells(start_row=1, start_column=1,
                          end_row=1, end_column=n_cols)

    # ── Row 2: "Room" banner over room columns ────────────────────────────────
    if room_cols:
        last_room_idx = first_room_idx + len(room_cols) - 1
        cell = worksheet.cell(row=2, column=first_room_idx, value="Room")
        cell.font      = Font(bold=True)
        cell.alignment = centre
        cell.fill      = room_fill
        worksheet.merge_cells(start_row=2, start_column=first_room_idx,
                              end_row=2, end_column=last_room_idx)

    # ── Row 3: headers ────────────────────────────────────────────────────────
    for idx, name in enumerate(headers, start=1):
        cell = worksheet.cell(row=3, column=idx, value=name)
        cell.font   = Font(bold=True, color="1F3864")
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        if idx == pwd_idx:
            cell.fill = pwd_fill
        elif idx == total_idx:
            cell.fill = total_fill
        elif idx >= first_room_idx:
            cell.fill = room_fill
        else:
            cell.fill = header_fill
    worksheet.row_dimensions[3].height = 36

    # ── Rows 4+: data ─────────────────────────────────────────────────────────
    last_data_row = len(pivot) + 3
    for r, (_, row) in enumerate(pivot.iterrows(), start=4):
        for idx, name in enumerate(headers, start=1):
            val = row.get(name)

            # Numeric columns: blank out zeros except Grand Total
            if name in num_cols:
                val = int(val or 0)
                if val == 0 and name != "Grand Total":
                    val = None

            cell = worksheet.cell(row=r, column=idx, value=val)
            cell.border = border

            if name == "Exam Date":
                cell.number_format = "DD-MMM-YYYY"
                cell.alignment = centre
            elif name in centre_always:
                cell.alignment = centre

            # Cell fills
            if idx == pwd_idx:
                cell.fill = pwd_fill
            elif idx == total_idx:
                cell.font = Font(bold=True)
                cell.fill = total_fill
            elif idx >= first_room_idx and val:
                cell.fill = room_fill   # only non-blank room cells tinted

    # ── Grand-total footer row ────────────────────────────────────────────────
    footer_row = last_data_row + 1
    # Merge only the pure text-label columns (before Grand Total)
    label_span_end = min(first_room_idx - 1, total_idx - 1, pwd_idx - 1) - 1
    label_span_end = max(label_span_end, 1)  # at minimum span col 1 only

    cell = worksheet.cell(row=footer_row, column=1, value="Grand Total")
    cell.font = Font(bold=True)
    if label_span_end > 1:
        worksheet.merge_cells(start_row=footer_row, start_column=1,
                              end_row=footer_row, end_column=label_span_end)
    # Style top-left of merge (others are MergedCell, read-only)
    c1 = worksheet.cell(row=footer_row, column=1)
    c1.fill   = footer_fill
    c1.border = border

    # Fill any unmerged label cols between label_span_end and first num col
    for fill_idx in range(label_span_end + 1, min(total_idx, pwd_idx)):
        fc = worksheet.cell(row=footer_row, column=fill_idx)
        fc.fill   = footer_fill
        fc.border = border

    for idx, name in enumerate(headers, start=1):
        if name not in num_cols:
            continue
        # Skip cells that fall inside the merged region (only col 1 is writable)
        if label_span_end > 1 and 2 <= idx <= label_span_end:
            continue
        letter = get_column_letter(idx)
        cell   = worksheet.cell(row=footer_row, column=idx)
        cell.value     = (f"=SUM({letter}4:{letter}{last_data_row})"
                          if len(pivot) else 0)
        cell.font      = Font(bold=True)
        cell.alignment = centre
        cell.border    = border
        if idx == pwd_idx:
            cell.fill = pwd_fill
        elif idx == total_idx:
            cell.fill = total_fill
        elif idx >= first_room_idx:
            cell.fill = room_fill
        else:
            cell.fill = footer_fill

    # ── Column widths ─────────────────────────────────────────────────────────
    for idx, name in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        if name in num_cols or name in ("Start Time", "End Time"):
            width = max(len(name) + 2, 10)
        elif name == "Exam Date":
            width = 14
        else:
            longest = max(
                [len(name)]
                + ([len(str(v)) for v in pivot[name].dropna().head(200)]
                   if name in pivot.columns else [])
            )
            width = min(max(longest + 2, 12), 45)
        worksheet.column_dimensions[letter].width = width

    # ── Freeze, filter, no gridlines ─────────────────────────────────────────
    # Freeze after the 3 date/time lead cols so they stay visible on scroll
    worksheet.freeze_panes = f"D4"
    if len(pivot):
        worksheet.auto_filter.ref = (
            f"A3:{get_column_letter(n_cols)}{last_data_row}"
        )
    worksheet.sheet_view.showGridLines = False


def to_excel(students_df, schedule_df, timetable_df, pivot_df=None, room_cols=None):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        if pivot_df is not None and not pivot_df.empty:
            sheet = writer.book.create_sheet("Pivot", 0)
            write_pivot_sheet(sheet, pivot_df, room_cols or [])
        students_df.to_excel(writer, sheet_name="Students", index=False)
        schedule_df.to_excel(writer, sheet_name="Exam Schedule", index=False)
        if not timetable_df.empty:
            timetable_df.to_excel(writer, sheet_name="Timetable", index=False)

        for name, frame in (
            ("Students", students_df),
            ("Exam Schedule", schedule_df),
            ("Timetable", timetable_df),
        ):
            if frame.empty:
                continue
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            for col_idx, col in enumerate(frame.columns, start=1):
                longest = max(
                    [len(str(col))] + [len(str(v)) for v in frame[col].head(200)]
                )
                sheet.column_dimensions[
                    sheet.cell(row=1, column=col_idx).column_letter
                ].width = min(max(longest + 2, 10), 55)
    buffer.seek(0)
    return buffer


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------

def show():
    st.title("📄Hall Ticket PDF → Excel")
    st.caption(
        "Upload one or more hall ticket PDFs (one student per page). "
        "Student details and the exam schedule are pulled out into a downloadable workbook."
    )

    uploaded = st.file_uploader(
        "Hall ticket PDFs", type="pdf", accept_multiple_files=True, key="ht_uploaded_files"
    )

    if not uploaded:
        st.info("Waiting for a PDF.")
        return

    all_students, all_schedule = [], []
    progress = st.progress(0.0, text="Reading PDFs…")

    for i, file in enumerate(uploaded, start=1):
        try:
            students_df, schedule_df = parse_pdf(file, source_name=file.name)
        except Exception as exc:  # keep one bad file from killing the batch
            st.error(f"Could not read {file.name}: {exc}")
            continue
        if students_df.empty:
            st.warning(f"No hall tickets found in {file.name}.")
        all_students.append(students_df)
        all_schedule.append(schedule_df)
        progress.progress(i / len(uploaded), text=f"Read {file.name}")

    progress.empty()

    students_df = pd.concat(all_students, ignore_index=True) if all_students else pd.DataFrame()
    schedule_df = pd.concat(all_schedule, ignore_index=True) if all_schedule else pd.DataFrame()

    if students_df.empty:
        st.error("Nothing could be extracted. Is this the same hall ticket layout?")
        return

    timetable_df = build_timetable(schedule_df)
    pivot_df, room_cols, _room_label_map = build_pivot(students_df, schedule_df)

    col1, col2, col3 = st.columns(3)
    col1.metric("Students", len(students_df))
    col2.metric("Exam rows", len(schedule_df))
    col3.metric("Distinct papers", len(timetable_df))

    missing = students_df[students_df["Subjects"] == 0]
    if not missing.empty:
        st.warning(
            f"{len(missing)} student(s) had no subject rows detected — check pages: "
            + ", ".join(str(p) for p in missing["Page"])
        )

    tab0, tab1, tab2, tab3 = st.tabs(
        ["Pivot", "Students", "Exam Schedule", "Timetable"]
    )
    with tab0:
        st.dataframe(pivot_df, use_container_width=True, hide_index=True)
    with tab1:
        st.dataframe(students_df, use_container_width=True, hide_index=True)
    with tab2:
        st.dataframe(schedule_df, use_container_width=True, hide_index=True)
    with tab3:
        st.dataframe(timetable_df, use_container_width=True, hide_index=True)

    st.download_button(
        "Download Excel workbook",
        data=to_excel(students_df, schedule_df, timetable_df, pivot_df, room_cols if room_cols else []),
        file_name="hall_tickets.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        key="ht_download_btn",
    )


main = show

if __name__ == "__main__":
    st.set_page_config(page_title="Hall Ticket Extractor", page_icon="📄", layout="wide")
    show()